"""Test script for the Excel export fixes (duplication + cluster composition)."""
import os
import tempfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from systemictau.exporter import export_to_excel

print("=== Testing Excel export fixes ===")

np.random.seed(123)
T = 60
# Different trajectories so we can detect copy bugs
tau_spatial_local = np.linspace(1.0, 1.8, T) + np.sin(np.linspace(0, 6, T)) * 0.1
tau_spatial_med   = tau_spatial_local * 0.92
tau_spatial_glob  = tau_spatial_med * 0.78

tau_mv_local = np.linspace(0.9, 1.6, T) + np.cos(np.linspace(0, 5, T)) * 0.15
tau_mv_med   = tau_mv_local * 0.88
tau_mv_glob  = tau_mv_med * 0.71   # deliberately different from med

def make_analysis(scale_taus, scale_recd, n_mods, med_groups=None, glob_groups=None, method_med="auto", method_glob="auto"):
    return {
        "scale_results": {
            "Local":  {"taus_global": scale_taus[0], "accum_T": np.cumsum(scale_taus[0]), "dtk": np.zeros(T), "t_star": 22},
            "Medium": {"taus_global": scale_taus[1], "accum_T": np.cumsum(scale_taus[1]), "dtk": np.zeros(T), "t_star": 31},
            "Global": {"taus_global": scale_taus[2], "accum_T": np.cumsum(scale_taus[2]), "dtk": np.zeros(T), "t_star": 45},
        },
        "scale_metrics": {
            "Local":  {"mean_tau": float(scale_taus[0].mean()), "tau_std": 0.07, "coherence": 0.68, "recd_T_final": scale_recd[0], "n_modules": n_mods[0]},
            "Medium": {"mean_tau": float(scale_taus[1].mean()), "tau_std": 0.05, "coherence": 0.81, "recd_T_final": scale_recd[1], "n_modules": n_mods[1]},
            "Global": {"mean_tau": float(scale_taus[2].mean()), "tau_std": 0.03, "coherence": 0.89, "recd_T_final": scale_recd[2], "n_modules": n_mods[2]},
        },
        "config": {
            "window_size": 13,
            "recd_enabled": True,
            "medium_groups": med_groups or {},
            "global_groups": glob_groups or {},
            "clustering_method": {"medium": method_med, "global": method_glob},
        }
    }

# === Case with manual Global clusters for BOTH perspectives ===
spatial_manual_glob = {"Macro_1": ["LocA", "LocB", "LocC"], "Macro_2": ["LocD", "LocE"], "Macro_3": ["LocF", "LocG"]}
mv_manual_glob = {"Macro_X": ["Var1", "Var2", "Var3"], "Macro_Y": ["Var4", "Var5", "Var6"], "Macro_Z": ["Var7"]}

spatial = make_analysis(
    [tau_spatial_local, tau_spatial_med, tau_spatial_glob],
    [150.0, 141.2, 128.7],
    [12, 4, 3],
    med_groups={"Cluster_1": ["LocA","LocB"], "Cluster_2": ["LocC","LocD"]},
    glob_groups=spatial_manual_glob,
    method_med="manual", method_glob="manual"
)

multivariate = make_analysis(
    [tau_mv_local, tau_mv_med, tau_mv_glob],
    [162.3, 149.8, 131.4],
    [7, 3, 3],
    med_groups={},
    glob_groups=mv_manual_glob,
    method_med="auto", method_glob="manual"
)

analyses = {"spatial": spatial, "multivariate": multivariate}
df_raw = pd.DataFrame({"week": range(90), "val": np.random.rand(90)})

print("Exporting with manual Global clusters on both perspectives...")
xlsx = export_to_excel(analyses, df_raw=df_raw, active_view="multivariate", window_size=13, include_raw_data=False)

with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
    tmp.write(xlsx)
    path = tmp.name

print("Written to", path)
wb = load_workbook(path)
print("Sheets:", wb.sheetnames)

# === Check Summary for duplication bug ===
print("\n--- Summary sheet (first data rows) ---")
ws = wb["Summary"]
for row in ws.iter_rows(min_row=6, max_row=12, values_only=True):
    print(row)

# Extract numeric values for multivariate Medium vs Global
summary_data = []
for row in ws.iter_rows(min_row=7, max_row=12, values_only=True):
    if row and row[0] == "multivariate":
        summary_data.append({"scale": row[1], "mean_tau": row[2], "recd": row[5], "tstar": row[6]})

print("\nMultivariate extracted from Summary:")
for d in summary_data:
    print("  ", d)

if len(summary_data) >= 2:
    med = next((x for x in summary_data if x["scale"] == "Medium"), None)
    glo = next((x for x in summary_data if x["scale"] == "Global"), None)
    if med and glo:
        same = (med["mean_tau"] == glo["mean_tau"]) and (med["recd"] == glo["recd"])
        print(f"\n>>> multivariate Medium vs Global are identical? {same}")
        assert not same, "BUG: multivariate Global duplicated Medium values!"
        print("PASS: Global values are different from Medium.")

# === Check Cluster_Composition ===
print("\n--- Cluster_Composition sheet ---")
ws_cc = wb["Cluster_Composition"]
for row in ws_cc.iter_rows(min_row=1, max_row=10, values_only=True):
    print(row)

cc_rows = list(ws_cc.iter_rows(min_row=2, values_only=True))
has_spatial_manual = any(r and r[0]=="spatial" and r[2]=="manual" for r in cc_rows if r)
has_mv_manual = any(r and r[0]=="multivariate" and r[2]=="manual" for r in cc_rows if r)
print(f"\nSpatial manual clusters present: {has_spatial_manual}")
print(f"Multivariate manual clusters present: {has_mv_manual}")
assert has_mv_manual, "BUG: multivariate manual Global clusters not recorded!"
print("PASS: Manual clusters recorded for multivariate Global.")

# === Check that all 6 rows are in Metrics_by_Scale ===
print("\n--- Metrics_by_Scale (checking 6 rows) ---")
ws_m = wb["Metrics_by_Scale"]
metrics_rows = [r for r in ws_m.iter_rows(min_row=2, values_only=True) if r and r[0] in ("spatial", "multivariate")]
print(f"Found {len(metrics_rows)} data rows for the two perspectives.")
assert len(metrics_rows) >= 6, "Should have at least 6 rows (2 perspectives x 3 scales)"
print("PASS: At least 6 scale rows present.")

wb.close()
os.unlink(path)
print("\n=== ALL CHECKS PASSED ===")